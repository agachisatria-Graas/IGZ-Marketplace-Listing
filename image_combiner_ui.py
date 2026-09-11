import io

import pandas as pd
import streamlit as st
import openpyxl
from openpyxl.styles import Font

import image_combiner


def render():
    st.title("Image Link Combiner")
    st.caption("Not connected to the marketplace listing tool — just a standalone helper.")
    st.write(
        "Combines images into one `\" ; \"`-separated string per SKU, in the "
        "right order — ready to paste into the listing tool's image columns."
    )

    mode = st.radio(
        "What do you want to do?",
        [
            "Combine Lazada/Shopee images + Zalora images (two files → one output with 2 tabs)",
            "Combine a single file (File Name + Image URL columns)",
            "Paste a list of URLs where the SKU_N pattern is IN the URL itself",
        ],
    )

    if mode.startswith("Combine Lazada/Shopee"):
        st.write(
            "Zalora needs images at a different dimension (**762 × 1100**) than "
            "Lazada/Shopee, so it needs its own separate source file. Upload "
            "both below — each is combined independently, and the result comes "
            "back as one Excel file with a **Lazada Images** tab and a "
            "**Zalora Images** tab."
        )
        col1, col2 = st.columns(2)
        with col1:
            st.subheader("Lazada / Shopee Images")
            lazada_combined, lazada_unmatched = _media_file_uploader("lazada")
        with col2:
            st.subheader("Zalora Images (762 × 1100)")
            zalora_combined, zalora_unmatched = _media_file_uploader("zalora")

        if lazada_combined or zalora_combined:
            st.divider()
            st.subheader("Combined output")
            buf = io.BytesIO()
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
            _write_combined_sheet(wb, "Lazada Images", lazada_combined or {})
            _write_combined_sheet(wb, "Zalora Images", zalora_combined or {})
            wb.save(buf)
            buf.seek(0)
            st.success(
                f"Lazada: {len(lazada_combined or {})} SKU(s) combined. "
                f"Zalora: {len(zalora_combined or {})} SKU(s) combined."
            )
            st.download_button(
                "Download lazada_zalora_images.xlsx",
                data=buf.read(),
                file_name="lazada_zalora_images.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            if lazada_unmatched:
                with st.expander(f"⚠️ Lazada: {len(lazada_unmatched)} File Name(s) didn't match SKU_N.ext"):
                    for u in lazada_unmatched:
                        st.write(u)
            if zalora_unmatched:
                with st.expander(f"⚠️ Zalora: {len(zalora_unmatched)} File Name(s) didn't match SKU_N.ext"):
                    for u in zalora_unmatched:
                        st.write(u)
        else:
            st.info("Upload at least one of the two files above to get started.")

    elif mode.startswith("Combine a single file"):
        st.write(
            "The order comes from the number in **File Name** (e.g. `_1`, `_2`, "
            "`_3`) — the **Image URL** column can be anything, even random "
            "upload links with no SKU in them at all."
        )
        combined, unmatched = _media_file_uploader("single")
        if combined is not None:
            _render_combined_result(combined, unmatched, unmatched_label="File Name(s)")
        else:
            st.info("Upload a file to get started.")

    else:
        pasted = st.text_area(
            "Paste image URLs, one per line",
            height=220,
            placeholder="https://example.com/DSM427123Y-M_1.jpg\n"
                        "https://example.com/DSM427123Y-M_2.jpg\n"
                        "https://example.com/HF-TUM-BUR_1.jpg",
        )
        if pasted.strip():
            urls = [line for line in pasted.splitlines() if line.strip()]
            combined, unmatched = image_combiner.group_and_combine(urls)
            _render_combined_result(combined, unmatched, unmatched_label="URL(s)")
        else:
            st.info("Paste some URLs to get started.")


def _media_file_uploader(key_prefix):
    """Renders a file uploader + column pickers, returns (combined, unmatched)
    or (None, None) if nothing's been uploaded yet."""
    uploaded_media_file = st.file_uploader(
        "Upload file (.xlsx or .csv)", type=["xlsx", "csv"], key=f"upload_{key_prefix}"
    )
    if uploaded_media_file is None:
        return None, None
    try:
        if uploaded_media_file.name.lower().endswith(".csv"):
            df_media = pd.read_csv(uploaded_media_file)
        else:
            df_media = pd.read_excel(uploaded_media_file)
    except Exception as e:
        st.error(f"Couldn't read that file: {e}")
        return None, None

    cols = list(df_media.columns)

    def guess(keywords, exclude=()):
        for c in cols:
            cl = str(c).lower()
            if any(k in cl for k in keywords) and not any(x in cl for x in exclude):
                return c
        return cols[0]

    guessed_filename_col = guess(["file", "name"], exclude=["date"])
    guessed_url_col = guess(["image", "url", "link"])

    filename_col = st.selectbox(
        "File Name column", cols, index=cols.index(guessed_filename_col),
        key=f"fname_{key_prefix}",
    )
    url_col = st.selectbox(
        "Image URL column", cols, index=cols.index(guessed_url_col),
        key=f"url_{key_prefix}",
    )
    st.dataframe(df_media[[filename_col, url_col]].head(5), use_container_width=True)

    pairs = list(zip(df_media[filename_col], df_media[url_col]))
    return image_combiner.group_and_combine_from_pairs(pairs)


def _write_combined_sheet(wb, sheet_title, combined):
    ws = wb.create_sheet(sheet_title)
    ws.cell(row=1, column=1, value="SKU").font = Font(bold=True)
    ws.cell(row=1, column=2, value="Combined Images").font = Font(bold=True)
    for i, (sku, joined) in enumerate(combined.items(), start=2):
        ws.cell(row=i, column=1, value=sku)
        ws.cell(row=i, column=2, value=joined)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 80
    return ws


def _render_combined_result(combined, unmatched, unmatched_label="item(s)"):
    if combined:
        st.success(f"Combined images for {len(combined)} SKU(s).")
        df_result = pd.DataFrame(
            [{"SKU": sku, "Combined Images": joined} for sku, joined in combined.items()]
        )
        st.dataframe(df_result, use_container_width=True, height=300)

        buf = io.BytesIO()
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        _write_combined_sheet(wb, "Combined Images", combined)
        wb.save(buf)
        buf.seek(0)
        st.download_button(
            "Download combined_images.xlsx",
            data=buf.read(),
            file_name="combined_images.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    else:
        st.warning("None of the rows matched the expected `SKU_N.ext` pattern.")

    if unmatched:
        with st.expander(f"⚠️ {len(unmatched)} {unmatched_label} didn't match the SKU_N.ext pattern"):
            for u in unmatched:
                st.write(u)
