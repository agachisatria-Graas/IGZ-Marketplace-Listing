"""Builds raw_data_template.xlsx and category_mapping_template.xlsx —
the two input files Agachi fills in."""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

RAW_COLUMNS = [
    ("parent_id", "Parent SKU", "Same value for every variant of one product. If the item has no variant, use the SKU itself.", "HF-DRYSTORAGE-HARBOR"),
    ("sku", "Seller SKU", "Unique seller SKU for this exact variant. REQUIRED, one row per SKU.", "DSM427123Y-M"),
    ("title", "Title", "Product title / name. This is what the Category Mapping file's 'Keyword in Title' matches against to auto-fill each marketplace's Category ID — always upload a Category Mapping file, since there's no manual category ID field here anymore.", "Hydro Flask Dry Storage Pouch Harbor"),
    ("main_description", "Main Description", "Plain text, one bullet/line per line. Combined with Long Description (line break in between) to form the final description used across all 4 marketplaces.", "Dry Storage locking technology slides to seal\nDurable coated fabric with welded seams"),
    ("long_description", "Long Description (optional)", "Plain text, one bullet/line per line. Appended after Main Description with a line break. Leave blank if you only need one description.", "Touchscreen window for device use\nIdeal size for phone, keys, wallet, sunscreen, snack"),
    ("description_script_override", "Description Script Override (optional)", "Paste ready-made HTML here to use for Lazada & Zalora instead of the auto-converted Main+Long Description. Leave blank to auto-convert.", ""),
    ("short_description", "Short Description (Lazada)", "One highlight per line. Becomes a bullet list on Lazada. Leave blank to auto-build from Lazada Item Specifications.", "Waterproof\nCompact size"),
    ("brand", "Brand", "Brand name.", "Hydro Flask"),
    ("variant_name_1", "Variant Name 1 (optional)", "e.g. Color. Leave blank if the item has no variants.", "Size"),
    ("variant_value_1", "Variant Value 1 (optional)", "e.g. Burgundy.", "M"),
    ("variant_name_2", "Variant Name 2 (optional)", "Second variant axis, e.g. Size.", ""),
    ("variant_value_2", "Variant Value 2 (optional)", "", ""),
    ("price", "Price", "Selling price, numbers only.", 849000),
    ("stock", "Stock / Quantity", "", 50),
    ("parent_images", "Parent Images", "Image URLs shared by all variants, separated by ' ; ' (space-semicolon-space).", "https://example.com/img1.jpg ; https://example.com/img2.jpg"),
    ("variant_images", "Variant Images (optional)", "Image URLs specific to this one variant only, separated by ' ; '. Leave blank if none. Used for Shopee/Lazada/TikTok — Zalora uses its own separate image field below.", ""),
    ("weight_kg", "Weight (kg)", "", 0.5),
    ("length_cm", "Length (cm)", "", 23),
    ("width_cm", "Width (cm)", "", 18.5),
    ("height_cm", "Height (cm)", "", 23),
    ("shopee_shipping_service", "Shopee Shipping Service", "Pasted as-is into Shopee's Shipping Service Details column.", "Reguler (Cashless):18000.00, Instant:18000.00"),
    ("shopee_item_specifications", "Shopee Item Specifications", "'Key=Value' pairs separated by ' ; '. Feeds Shopee's Product Specification columns only.", "Brand=Hydro Flask ; Material=Stainless Steel"),
    ("lazada_item_specifications", "Lazada Item Specifications", "'Key=Value' pairs separated by ' ; '. Feeds Lazada's Product Specification columns, and its Short Description bullets if that field is left blank.", "Brand=Hydro Flask ; Material=Stainless Steel"),
    ("tiktok_item_specifications", "TikTok Item Specifications", "'Key=Value' pairs separated by ' ; '. Written to generic Specification columns in the output — TikTok's real template needs these under specific attribute-ID columns per category, so copy them over manually before uploading.", "Material=Stainless Steel"),
    ("zalora_gender", "Zalora Gender", "Also used as the 'Gender' to match against your Category Mapping file, across all 4 marketplaces — not just Zalora.", "Unisex"),
    ("zalora_subcat_type", "Zalora Sub Cat Type", "", "Sports Metal Water Bottles"),
    ("zalora_color_family", "Zalora Color Family", "", "grey"),
    ("zalora_color", "Zalora Color", "", "Harbor"),
    ("zalora_images", "Zalora Images", "Zalora's own image set, separated by ' ; '. Zalora requires different image dimensions/crops than the other 3 marketplaces, so give it its own set here rather than reusing Parent/Variant Images.", "https://example.com/zalora1.jpg ; https://example.com/zalora2.jpg"),
]

# Category mapping template: one sheet per marketplace, each with
# (Gender, Keyword in Title, Category ID) — matching Agachi's real file layout.
CATEGORY_SHEETS = {
    "Shopee Category": ("Gender", "Keyword in Title", "ID Shopee Category id", "Unisex", "Tumbler", 101692),
    "Lazada category": ("Gender", "Keyword in Title", "ID Lazada Category id", "Unisex", "Tumbler", 42054601),
    "Tiktok Category": ("Gender", "Keyword in Title", "ID Tiktok Category id", "Unisex", "Tumbler", "Perlengkapan Minum/Botol Air"),
    "Zalora Category": ("Gender", "Keyword in Title", "ID Zalora Category id", "Unisex", "Tumbler", "12352 - Sports / Sports Pria / Perlengkapan Olahraga / Lifestyle / Water Bottles"),
}

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5496")
NOTE_FONT = Font(name="Arial", italic=True, size=9, color="666666")
EXAMPLE_FONT = Font(name="Arial", italic=True, size=10, color="1F7A1F")
INPUT_FILL = PatternFill("solid", fgColor="FFFFCC")


def build_raw_data_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw Data"
    for i, (key, label, note, example) in enumerate(RAW_COLUMNS, start=1):
        col = get_column_letter(i)
        c1 = ws.cell(row=1, column=i, value=label)
        c1.font = HEADER_FONT
        c1.fill = HEADER_FILL
        c1.alignment = Alignment(wrap_text=True, vertical="center")
        c2 = ws.cell(row=2, column=i, value=note)
        c2.font = NOTE_FONT
        c2.alignment = Alignment(wrap_text=True, vertical="top")
        c3 = ws.cell(row=3, column=i, value=example)
        c3.font = EXAMPLE_FONT
        c3.fill = INPUT_FILL
        ws.column_dimensions[col].width = 26
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 60
    ws.freeze_panes = "A4"
    return wb


def build_category_mapping_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name, (h1, h2, h3, g, kw, cid) in CATEGORY_SHEETS.items():
        ws = wb.create_sheet(sheet_name)
        for i, h in enumerate((h1, h2, h3), start=1):
            c = ws.cell(row=1, column=i, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
        ws.cell(row=2, column=1, value=g)
        ws.cell(row=2, column=2, value=kw)
        ws.cell(row=2, column=3, value=cid)
        for i in range(1, 4):
            ws.column_dimensions[get_column_letter(i)].width = 30
        ws.freeze_panes = "A2"
    return wb


if __name__ == "__main__":
    build_raw_data_workbook().save("/home/claude/listing_tool/raw_data_template.xlsx")
    build_category_mapping_workbook().save(
        "/home/claude/listing_tool/category_mapping_template.xlsx"
    )
    print("saved both templates")
