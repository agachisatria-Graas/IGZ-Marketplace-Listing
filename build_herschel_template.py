"""Builds herschel_raw_data_template.xlsx and
herschel_category_mapping_template.xlsx — the two input files for the
Herschel Marketplace Listing Tool. Kept separate from build_template.py
(Hydro Flask's) since this is a fully separate tool."""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HERSCHEL_RAW_COLUMNS = [
    ("parent_id", "Parent SKU", "Same value for every variant of one product. If the item has no variant, use the SKU itself.", "HS-LITTLEAMERICA-BLK"),
    ("sku", "Seller SKU", "Unique seller SKU for this exact variant. REQUIRED, one row per SKU.", "2100326232"),
    ("title", "Title", "Product title / name.", "Herschel Little America Backpack"),
    ("main_description", "Main Description", "Plain text, one bullet/line per line.", "Signature vegan leather bottom accents\nAdjustable, padded shoulder straps"),
    ("main_description_2", "Main Description 2", "Plain text, one bullet/line per line. Joined after Main Description with a line break.", "Internal media pocket with headphone port\nFront zippered pocket"),
    ("measurement", "Measurement", "Plain text — dimensions/capacity etc. Joined after Main Description 2 with a line break.", "H 48cm x W 31cm x D 20cm | 25L"),
    ("description_script_override", "Description Script Override (optional)", "Paste ready-made HTML here to use for Lazada & Zalora instead of the auto-converted Main+Main2+Measurement. Leave blank to auto-convert.", ""),
    ("short_description", "Short Description (Lazada)", "One highlight per line. Becomes a bullet list on Lazada. Leave blank to auto-build from Lazada Item Specifications.", "Vegan leather details\nPadded laptop sleeve"),
    ("brand", "Brand", "Brand name.", "Herschel"),
    ("variant_name_1", "Variant Name 1 (optional)", "e.g. Color. Leave blank if the item has no variants.", "Color"),
    ("variant_value_1", "Variant Value 1 (optional)", "e.g. Black.", "Black"),
    ("variant_name_2", "Variant Name 2 (optional)", "Second variant axis.", ""),
    ("variant_value_2", "Variant Value 2 (optional)", "", ""),
    ("price", "Price", "Selling price, numbers only.", 1299000),
    ("parent_images", "Parent Images", "Image URLs shared by all variants, separated by ' ; ' (space-semicolon-space). Used for Shopee/Lazada/TikTok.", "https://example.com/img1.jpg ; https://example.com/img2.jpg"),
    ("variant_images", "Variant Images (optional)", "Image URLs specific to this one variant only, separated by ' ; '. Used for Shopee/Lazada/TikTok — Zalora uses its own separate image field below.", ""),
    ("weight_kg", "Weight (kg)", "If copying from a spec sheet showing both units (e.g. '1.10 lb / 0.5'), paste it as-is — the tool automatically takes the number AFTER the '/' as the kg value. A plain number with no slash also works fine.", "1.10 lb / 0.5"),
    ("length_cm", "Length (cm)", "", 48),
    ("width_cm", "Width (cm)", "", 31),
    ("height_cm", "Height (cm)", "", 20),
    ("product_type", "Product Type", "e.g. BAGS or ACCESSORY. Used with Specific Category and Gender for an EXACT match against your Category Mapping file — not a keyword search.", "BAGS"),
    ("specific_category", "Specific Category", "e.g. BACKPACKS, TOTES, MESSENGER, DUFFLES, WAIST BAGS, LUNCH BOX, ORGANIZER.", "BACKPACKS"),
    ("gender", "Gender*", "Code: US, WN, or UK. Also translated into Zalora's own Gender field (US=Men, WN=Women, UK=Kids).", "US"),
    ("material", "Material", "Free text, one point per line (paste your full bullet list here if you like) — the tool automatically picks out the first recognizable material keyword (e.g. Polyester, Leather, Nylon) to use in Shopee/Lazada/TikTok/Zalora's Material fields.", "Custom print inspired by the worlds of Minecraft\n100% recycled 600D polyester, excluding trims"),
    ("shopee_shipping_service", "Shopee Shipping Service", "Pasted as-is into Shopee's Shipping Service Details column.", "Reguler (Cashless):18000.00, Instant:18000.00"),
    ("shopee_item_specifications", "Shopee Item Specifications", "'Key=Value' pairs separated by ' ; ', appended AFTER the automatic Brand=Herschel and Material=... entries.", "Capacity=25L"),
    ("lazada_item_specifications", "Lazada Item Specifications", "'Key=Value' pairs separated by ' ; ', appended AFTER the automatic delivery/Hazmat/material entries. Also feeds Lazada's Short Description bullets if that field is left blank.", "Capacity=25L"),
    ("tiktok_item_specifications", "TikTok Item Specifications", "'Key=Value' pairs separated by ' ; '. A key matching one of TikTok's named columns (Jenis Kulit, Pola, Acara, Gaya, Instruksi Mencuci, Tipe Pengencang, Tipe Tas, Fitur, Bahan) lands directly in that column — e.g. 'Pola=Solid' fills the Pola column. Unmatched keys are dropped.", "Pola=Solid"),
    ("season", "Season", "For Zalora's Season column.", "Autumn-Winter"),
    ("year", "Year", "For Zalora's Year column.", 2026),
    ("zalora_color", "Zalora Color", "This also drives automatic ColorFamily classification — the tool looks at the first Zalora Image's dominant color first, falling back to matching keywords in this Color name if no image is available.", "Pink Sheep"),
    ("zalora_images", "Zalora Images", "Zalora's own image set (different dimensions than the other 3 marketplaces), separated by ' ; '.", "https://example.com/zalora1.jpg ; https://example.com/zalora2.jpg"),
]

# Category mapping template: one sheet per marketplace, each with
# (Product Type, Specific Category, Gender*, Category ID) — matching
# Agachi's real IGZ_Herschel_category_sheet.xlsx layout exactly.
HERSCHEL_CATEGORY_SHEETS = {
    "Shopee Category": ("Product Type", "Specific Category", "Gender*", "ID Shopee Category id", "BAGS", "BACKPACKS", "US", 100564),
    "Lazada category": ("Product Type", "Specific Category", "Gender*", "ID Lazada Category id", "BAGS", "BACKPACKS", "US", 10100762),
    "Tiktok Category": ("Product Type", "Specific Category", "Gender*", "ID Tiktok Category id", "BAGS", "BACKPACKS", "US", "Tas Fungsional/Ransel"),
    "Zalora Category": ("Product Type", "Specific Category", "Gender*", "ID Zalora Category id", "BAGS", "BACKPACKS", "US", "317 - Tas / Tas Pria / Backpack"),
}

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="7030A0")  # purple, to visually distinguish from Hydro Flask's blue
NOTE_FONT = Font(name="Arial", italic=True, size=9, color="666666")
EXAMPLE_FONT = Font(name="Arial", italic=True, size=10, color="1F7A1F")
INPUT_FILL = PatternFill("solid", fgColor="FFFFCC")


def build_herschel_raw_data_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw Data"
    for i, (key, label, note, example) in enumerate(HERSCHEL_RAW_COLUMNS, start=1):
        col = get_column_letter(i)
        c1 = ws.cell(row=1, column=i, value=label)
        c1.font = HEADER_FONT
        c1.fill = HEADER_FILL
        c1.alignment = Alignment(wrap_text=True, vertical="center")
        c2 = ws.cell(row=2, column=i, value=note)
        c2.font = NOTE_FONT
        c2.alignment = Alignment(wrap_text=True, vertical="top")
        c3 = ws.cell(row=3, column=i, value=example)
        c3.font = EXAMPLE_FONT
        c3.fill = INPUT_FILL
        ws.column_dimensions[col].width = 26
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 70
    ws.freeze_panes = "A4"
    return wb


def build_herschel_category_mapping_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name, (h1, h2, h3, h4, pt, sc, g, cid) in HERSCHEL_CATEGORY_SHEETS.items():
        ws = wb.create_sheet(sheet_name)
        for i, h in enumerate((h1, h2, h3, h4), start=1):
            c = ws.cell(row=1, column=i, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
        ws.cell(row=2, column=1, value=pt)
        ws.cell(row=2, column=2, value=sc)
        ws.cell(row=2, column=3, value=g)
        ws.cell(row=2, column=4, value=cid)
        for i in range(1, 5):
            ws.column_dimensions[get_column_letter(i)].width = 28
        ws.freeze_panes = "A2"
    return wb


if __name__ == "__main__":
    build_herschel_raw_data_workbook().save(
        "/home/claude/listing_tool/herschel_raw_data_template.xlsx"
    )
    build_herschel_category_mapping_workbook().save(
        "/home/claude/listing_tool/herschel_category_mapping_template.xlsx"
    )
    print("saved both Herschel templates")
