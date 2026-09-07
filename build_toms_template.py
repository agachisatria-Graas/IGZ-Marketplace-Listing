"""Builds toms_raw_data_template.xlsx and
toms_category_mapping_template.xlsx — the two input files for the Toms
Marketplace Listing Tool. Kept separate from the other two brands' template
builders since this is a fully separate tool."""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

TOMS_RAW_COLUMNS = [
    ("parent_id", "Parent SKU", "Same value for every variant of one product. If the item has no variant, use the SKU itself.", "2100333242"),
    ("sku", "Seller SKU", "Unique seller SKU for this exact variant. REQUIRED, one row per SKU.", "2100333242-5.5"),
    ("title", "Title", "Product title / name. Also scanned for BOTH the 'Gender in Title' and 'Words in Title' keywords in your Category Mapping file — e.g. must contain 'Women' AND 'Mule' to match a Women+Mule mapping row.", "Women Alpargata Classic - Black/White"),
    ("main_description", "Main Description", "Plain text, one bullet/line per line.", "The TOMS signature Alpargata slip-on\nHeritage Canvas upper"),
    ("main_description_2", "Main Description 2", "Plain text, one bullet/line per line. Joined after Main Description with a line break.", "Injected TPR outsole for flexibility"),
    ("description_script_override", "Description Script Override (optional)", "Paste ready-made HTML here to use for Lazada & Zalora instead of the auto-converted Main+Main2. Leave blank to auto-convert.", ""),
    ("short_description", "Short Description (Lazada)", "One highlight per line. Becomes a bullet list on Lazada. Leave blank to auto-build from Lazada Item Specifications.", "Slip-on style\nCanvas upper"),
    ("brand", "Brand", "Brand name.", "Toms"),
    ("variant_name_1", "Variant Name 1 (optional)", "e.g. Size.", "Size"),
    ("variant_value_1", "Variant Value 1 (optional)", "e.g. US 5.5.", "US 5.5"),
    ("variant_name_2", "Variant Name 2 (optional)", "Second variant axis.", ""),
    ("variant_value_2", "Variant Value 2 (optional)", "", ""),
    ("price", "Price", "Selling price, numbers only.", 899000),
    ("stock", "Stock / Quantity", "", 20),
    ("parent_images", "Parent Images", "Image URLs shared by all variants, separated by ' ; ' (space-semicolon-space). Used for Shopee/Lazada/TikTok.", "https://example.com/img1.jpg ; https://example.com/img2.jpg"),
    ("variant_images", "Variant Images (optional)", "Image URLs specific to this one variant only, separated by ' ; '. Used for Shopee/Lazada/TikTok — Zalora uses its own separate image field below.", ""),
    ("weight_kg", "Weight (kg)", "", 1),
    ("length_cm", "Length (cm)", "", 40),
    ("width_cm", "Width (cm)", "", 23),
    ("height_cm", "Height (cm)", "", 20),
    ("specific_category", "Specific Category", "e.g. SNEAKERS, BCKSANDALS, SLIPONS. Used together with the Title keywords for an EXACT match against your Category Mapping file's Specific Category column.", "SNEAKERS"),
    ("gender", "Gender*", "Code, e.g. WN. Translated into Zalora's own Gender field (WN=Female confirmed so far — more codes to be added as they come up).", "WN"),
    ("material", "Material (optional)", "Free text — passed straight into Zalora's Material field as-is (not auto-extracted like Herschel's). Shopee/Lazada's Material spec is always the fixed 'Rubber' regardless of this field.", "Canvas upper, rubber outsole"),
    ("shopee_shipping_service", "Shopee Shipping Service", "Pasted as-is into Shopee's Shipping Service Details column.", "Reguler (Cashless):18000.00, Instant:18000.00"),
    ("shopee_item_specifications", "Shopee Item Specifications", "'Key=Value' pairs separated by ' ; ', appended AFTER the automatic Brand=Toms and Material=Rubber entries.", "Toe Style=Round Toe"),
    ("lazada_item_specifications", "Lazada Item Specifications", "'Key=Value' pairs separated by ' ; ', appended AFTER the automatic delivery/Hazmat/material entries. Also feeds Lazada's Short Description bullets if that field is left blank.", "Toe Style=Round Toe"),
    ("tiktok_item_specifications", "TikTok Item Specifications", "'Key=Value' pairs separated by ' ; '. A key matching one of TikTok's named columns (Bahan, Acara, Musim, Bentuk Jari Kaki, Tinggi Hak, Tipe Pengencang, Tipe Hak, Sertifikat SNI) lands directly in that column. Bahan defaults to 'Rubber' and Musim defaults to the Season field below unless overridden here. Unmatched keys are dropped.", "Bentuk Jari Kaki=Round"),
    ("season", "Season", "For Zalora's Season column, and TikTok's Musim column (unless overridden in TikTok Item Specifications).", "AUTUMN WINTER"),
    ("year", "Year", "For Zalora's Year column.", 2026),
    ("zalora_color", "Zalora Color", "This also drives automatic ColorFamily classification — the tool looks at the first Zalora Image's dominant color first, falling back to matching keywords in this Color name if no image is available.", "Black/White"),
    ("zalora_images", "Zalora Images", "Zalora's own image set (different dimensions than the other 3 marketplaces), separated by ' ; '.", "https://example.com/zalora1.jpg ; https://example.com/zalora2.jpg"),
]

# Category mapping template: one sheet per marketplace, each with
# (Gender in Title, Words in Title, Specific Category, Category ID) —
# matching Agachi's real IGZ_Toms_category_sheet.xlsx layout exactly.
TOMS_CATEGORY_SHEETS = {
    "Shopee Category": ("Gender in Title", "Words in Title", "Specific Category", "ID Shopee Category id", "Women", "Mary Jane", "SNEAKERS", 100591),
    "Lazada category": ("Gender in Title", "Words in Title", "Specific Category", "ID Lazada Category id", "Women", "Mary Jane", "SNEAKERS", 14883),
    "Tiktok Category": ("Gender in Title", "Words in Title", "Specific Category", "ID Tiktok Category id", "Women", "Mary Jane", "SNEAKERS", "Sepatu Wanita/Sepatu Mary Jane"),
    "Zalora Category": ("Gender in Title", "Words in Title", "Specific Category", "ID Zalora Category id", "Women", "Mary Jane", "SNEAKERS", "2171 - Sepatu / Sepatu Wanita / Slip On"),
}

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="C00000")  # dark red, to visually distinguish from the other two tools
NOTE_FONT = Font(name="Arial", italic=True, size=9, color="666666")
EXAMPLE_FONT = Font(name="Arial", italic=True, size=10, color="1F7A1F")
INPUT_FILL = PatternFill("solid", fgColor="FFFFCC")


def build_toms_raw_data_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw Data"
    for i, (key, label, note, example) in enumerate(TOMS_RAW_COLUMNS, start=1):
        col = get_column_letter(i)
        c1 = ws.cell(row=1, column=i, value=label)
        c1.font = HEADER_FONT
        c1.fill = HEADER_FILL
        c1.alignment = Alignment(wrap_text=True, vertical="center")
        c2 = ws.cell(row=2, column=i, value=note)
        c2.font = NOTE_FONT
        c2.alignment = Alignment(wrap_text=True, vertical="top")
        c3 = ws.cell(row=3, column=i, value=example)
        c3.font = EXAMPLE_FONT
        c3.fill = INPUT_FILL
        ws.column_dimensions[col].width = 26
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 70
    ws.freeze_panes = "A4"
    return wb


def build_toms_category_mapping_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name, (h1, h2, h3, h4, gw, w, sc, cid) in TOMS_CATEGORY_SHEETS.items():
        ws = wb.create_sheet(sheet_name)
        for i, h in enumerate((h1, h2, h3, h4), start=1):
            c = ws.cell(row=1, column=i, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
        ws.cell(row=2, column=1, value=gw)
        ws.cell(row=2, column=2, value=w)
        ws.cell(row=2, column=3, value=sc)
        ws.cell(row=2, column=4, value=cid)
        for i in range(1, 5):
            ws.column_dimensions[get_column_letter(i)].width = 28
        ws.freeze_panes = "A2"
    return wb


if __name__ == "__main__":
    build_toms_raw_data_workbook().save(
        "/home/claude/listing_tool/toms_raw_data_template.xlsx"
    )
    build_toms_category_mapping_workbook().save(
        "/home/claude/listing_tool/toms_category_mapping_template.xlsx"
    )
    print("saved both Toms templates")
