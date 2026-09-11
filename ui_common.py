"""
Shared utilities used by every brand's UI module (hydro_flask_ui.py,
herschel_ui.py, toms_ui.py, and future brand UI files).

Keeping this separate from app.py means app.py itself stays a thin
orchestrator — just imports + tab layout — no matter how many brands get
added.
"""
import io

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

PLATFORM_LABELS = {
    "shopee": "Shopee", "lazada": "Lazada", "tiktok": "TikTok Shop", "zalora": "Zalora",
}

# Used by every brand's category-mapping-workbook loader to match a sheet
# name (e.g. "Shopee Category", "Lazada category") to a platform key.
CATEGORY_SHEET_PLATFORM = {
    "shopee": "shopee", "lazada": "lazada", "tiktok": "tiktok", "zalora": "zalora",
}


def file_bytes(path):
    with open(path, "rb") as f:
        return f.read()


def rows_to_xlsx_bytes(platform, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
